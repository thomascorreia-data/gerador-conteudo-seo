"""
Envia as descrições geradas (empresas_descricoes.xlsx) pra API de páginas de
empresa da Buser (produção): cria (POST) ou atualiza (PATCH, se já existir)
a company-page de cada empresa.

Mapeamento planilha -> API:
    Nome da Empresa       -> company_slug (sem acento, minúsculo, espaço -> "-")
    Descrição Vendas      -> summary_text
    Descrição Informativo -> about_text

Uso:
    ./venv/bin/python enviando--buser.py --dry-run          # só mostra o que seria enviado
    ./venv/bin/python enviando--buser.py --dry-run --limite 3
    ./venv/bin/python enviando--buser.py --limite 1          # envia só a 1a empresa, de verdade
    ./venv/bin/python enviando--buser.py                     # envia todas, de verdade
"""

import argparse
import os
import re
import sys
import time
import unicodedata

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

BASE_URL = "https://www.buser.com.br"
ENDPOINT = f"{BASE_URL}/api/pages/integration/company-pages"
XLSX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "empresas_descricoes.xlsx")

API_KEY = os.getenv("BUSER_API_KEY")


def slugify(nome: str) -> str:
    """Mesma receita usada em formatos/transformacaoJson.py: sem acento,
    minúsculo, qualquer sequência de não-alfanumérico vira um único '-'."""
    texto = unicodedata.normalize("NFKD", nome)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.lower().strip()
    texto = re.sub(r"[^a-z0-9]+", "-", texto).strip("-")
    return texto


def carregar_empresas(caminho: str) -> list:
    wb = load_workbook(caminho, data_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    _cabecalho, *dados = linhas

    empresas = []
    for nome, informativo, vendas in dados:
        if not nome:
            continue
        empresas.append({
            "nome": nome,
            "company_slug": slugify(nome),
            "summary_text": (vendas or "").strip(),
            "about_text": (informativo or "").strip(),
        })
    return empresas


def _headers():
    return {"X-API-KEY": API_KEY, "Content-Type": "application/json"}


def enviar(empresa: dict, dry_run: bool) -> None:
    slug = empresa["company_slug"]
    payload = {
        "company_slug": slug,
        "summary_text": empresa["summary_text"],
        "contact_text": " ",
        "contact_email": " ",
        "about_text": empresa["about_text"],
        "faqs": [],
    }

    if dry_run:
        print(f"[DRY-RUN] {empresa['nome']!r} -> slug={slug!r}")
        print(f"          summary_text: {payload['summary_text'][:80]}...")
        print(f"          about_text:   {payload['about_text'][:80]}...")
        return

    resposta = requests.post(ENDPOINT, headers=_headers(), json=payload, timeout=30)

    if resposta.status_code == 409:
        # Empresa já tem página: cai pro PATCH (update parcial) em vez de falhar.
        resposta = requests.patch(
            f"{ENDPOINT}/{slug}",
            headers=_headers(),
            json={"summary_text": payload["summary_text"], "about_text": payload["about_text"]},
            timeout=30,
        )
        acao = "PATCH (já existia)"
    else:
        acao = "POST"

    print(f"{empresa['nome']} ({slug}) [{acao}]: {resposta.status_code}")
    if resposta.status_code >= 400:
        print(f"  -> {resposta.text[:300]}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="só mostra o que seria enviado, não chama a API")
    parser.add_argument("--limite", type=int, default=None, help="processa só as N primeiras empresas (pra teste)")
    args = parser.parse_args()

    if not args.dry_run and not API_KEY:
        sys.exit("BUSER_API_KEY não configurada no .env")

    empresas = carregar_empresas(XLSX_PATH)
    if args.limite:
        empresas = empresas[: args.limite]

    print(f"{len(empresas)} empresa(s) a processar ({'dry-run' if args.dry_run else 'ENVIO REAL para ' + BASE_URL}).\n")
    for empresa in empresas:
        enviar(empresa, dry_run=args.dry_run)
        if not args.dry_run:
            time.sleep(0.5)


if __name__ == "__main__":
    main()
